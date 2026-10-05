import asyncio
import hashlib
import os
import re
from io import BytesIO
from zipfile import BadZipFile, ZipFile

import fitz
import pytesseract
from docx import Document
from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import StreamingResponse
from PIL import Image
from pydantic import BaseModel
from pypdf import PdfReader

app = FastAPI(
    title="PlacePrep AI Service",
    version="0.1.0",
    description="Service boundary for resume parsing, explainable scoring, and AI adapters.",
)


MAX_RESUME_BYTES = 5 * 1024 * 1024
EMAIL_PATTERN = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.IGNORECASE)
PHONE_PATTERN = re.compile(r"(?:\+?\d[\d\s().-]{7,}\d)")
SKILL_TERMS = {
    "languages": ["python", "java", "javascript", "typescript", "c++", "c#", "sql"],
    "frameworks": ["react", "node.js", "express", "django", "fastapi", "spring", "angular"],
    "tools": ["git", "docker", "kubernetes", "excel", "power bi", "tableau", "linux"],
    "databases": ["postgresql", "mysql", "mongodb", "redis", "sqlite"],
    "softSkills": ["communication", "leadership", "teamwork", "collaboration", "problem solving"],
}
SECTION_ALIASES = {
    "summary": {"summary", "profile", "objective", "professional summary"},
    "education": {"education", "academic background", "qualifications"},
    "skills": {"skills", "technical skills", "technologies"},
    "projects": {"projects", "academic projects", "personal projects"},
    "experience": {"experience", "work experience", "internships", "employment"},
    "certifications": {"certifications", "certificates"},
    "achievements": {"achievements", "awards", "honors"},
}


def extract_docx(payload: bytes) -> str:
    try:
        with ZipFile(BytesIO(payload)) as archive:
            members = archive.infolist()
            expanded_size = sum(member.file_size for member in members)
            if len(members) > 500 or expanded_size > 30 * 1024 * 1024 or "word/document.xml" not in archive.namelist():
                raise HTTPException(status_code=422, detail={"code": "INVALID_DOCX", "message": "The DOCX file has an invalid or unsafe document structure."})
        document = Document(BytesIO(payload))
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(status_code=422, detail={"code": "INVALID_DOCX", "message": "The DOCX file is damaged or unreadable."}) from error
    content = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        content.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
    return "\n".join(content)


def extract_pdf(payload: bytes) -> tuple[str, int, bool]:
    try:
        reader = PdfReader(BytesIO(payload), strict=False)
        if reader.is_encrypted:
            try:
                decrypted = reader.decrypt("")
            except Exception as error:
                raise HTTPException(status_code=422, detail={"code": "PASSWORD_PROTECTED", "message": "This PDF is password-protected. Remove the password and upload it again."}) from error
            if decrypted == 0:
                raise HTTPException(status_code=422, detail={"code": "PASSWORD_PROTECTED", "message": "This PDF is password-protected. Remove the password and upload it again."})
        page_count = len(reader.pages)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(status_code=422, detail={"code": "INVALID_PDF", "message": "The PDF is damaged or cannot be read."}) from error

    did_ocr = False
    if len(text.strip()) < 100:
        did_ocr = True
        try:
            document = fitz.open(stream=payload, filetype="pdf")
            pages = []
            for page in document[: min(page_count, 5)]:
                pixmap = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
                image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
                pages.append(pytesseract.image_to_string(image, lang="eng+hin"))
            text = "\n".join(pages)
        except Exception:
            text = ""
    return text, page_count, did_ocr


def parse_sections(text: str) -> dict[str, str]:
    sections: dict[str, list[str]] = {key: [] for key in SECTION_ALIASES}
    active_section: str | None = None
    for line in text.splitlines():
        cleaned = re.sub(r"[#:|]+$", "", line.strip()).strip()
        normalized = re.sub(r"[^a-z ]", "", cleaned.lower()).strip()
        matching_section = next((key for key, aliases in SECTION_ALIASES.items() if normalized in aliases), None)
        if matching_section:
            active_section = matching_section
        elif active_section and cleaned:
            sections[active_section].append(cleaned)
    return {key: "\n".join(lines) for key, lines in sections.items()}


def parse_resume(text: str, page_count: int, used_ocr: bool) -> dict:
    sections = parse_sections(text)
    lowered = text.lower()
    found_skills = {
        group: [term for term in terms if term in lowered]
        for group, terms in SKILL_TERMS.items()
    }
    email_match = EMAIL_PATTERN.search(text)
    phone_match = PHONE_PATTERN.search(text)
    project_lines = [line.strip(" -*\t") for line in sections["projects"].splitlines() if line.strip()]
    experience_lines = [line.strip(" -*\t") for line in sections["experience"].splitlines() if line.strip()]
    warnings = []
    if not text.strip():
        warnings.append("Resume cannot be read by ATS.")
    if used_ocr:
        warnings.append("OCR was used. Review extracted content for accuracy.")
    if page_count > 3:
        warnings.append("Resume exceeds three pages.")
    if not email_match:
        warnings.append("No email address was detected.")

    return {
        "contact": {
            "email": email_match.group(0) if email_match else None,
            "phone": phone_match.group(0) if phone_match else None,
        },
        "summary": sections["summary"],
        "education": sections["education"],
        "skills": found_skills,
        "projects": project_lines,
        "experience": experience_lines,
        "certifications": sections["certifications"].splitlines(),
        "achievements": sections["achievements"].splitlines(),
        "rawSections": sections,
        "metadata": {
            "pageCount": page_count,
            "characterCount": len(text),
            "usedOcr": used_ocr,
            "readable": len(text.strip()) >= 100,
            "warnings": warnings,
        },
    }


# ── Email / Cover Letter Generator ──────────────────────────────────────────

class EmailGenerateRequest(BaseModel):
    candidateName: str
    candidateEmail: str | None = None
    skills: list[str] = []
    experience: list[str] = []
    projects: list[str] = []
    education: str = ""
    jobTitle: str
    companyName: str
    jobDescription: str = ""
    requiredSkills: list[str] = []
    type: str = "cover_letter"


def build_email_text(req: EmailGenerateRequest) -> str:
    skills_str = ", ".join(req.skills[:12]) or "various technical skills"
    projects_str = "; ".join(req.projects[:3]) or "several academic and personal projects"
    experience_str = req.experience[0] if req.experience else ""
    match_skills = [s for s in req.requiredSkills if s.lower() in skills_str.lower()][:5]
    match_str = ", ".join(match_skills) if match_skills else skills_str.split(",")[0].strip()

    if req.type == "cold_email":
        return (
            f"Subject: Application for {req.jobTitle} at {req.companyName}\n\n"
            f"Dear Hiring Team at {req.companyName},\n\n"
            f"I am {req.candidateName}, a motivated candidate with hands-on experience in {skills_str}. "
            f"I came across the {req.jobTitle} opening and believe my background is a strong match for what your team is building.\n\n"
            f"{'I have worked on ' + experience_str + ', which gave me direct exposure to the challenges your team likely faces. ' if experience_str else ''}"
            f"My projects include {projects_str}, demonstrating my ability to deliver end-to-end solutions. "
            f"I am particularly aligned with your requirement for {match_str}.\n\n"
            f"I would love the opportunity to contribute to {req.companyName}. "
            f"Could we schedule a brief call this week?\n\n"
            f"Best regards,\n{req.candidateName}"
            f"{chr(10) + req.candidateEmail if req.candidateEmail else ''}"
        )
    return (
        f"{req.candidateName}\n"
        f"{req.candidateEmail or ''}\n\n"
        f"Hiring Manager\n{req.companyName}\n\n"
        f"Dear Hiring Manager,\n\n"
        f"I am writing to express my strong interest in the {req.jobTitle} position at {req.companyName}. "
        f"With a solid foundation in {skills_str} and a track record of delivering impactful work, "
        f"I am confident I can contribute meaningfully to your team.\n\n"
        f"{'In my previous role, ' + experience_str + ' ' if experience_str else ''}"
        f"I have built and shipped {projects_str}. "
        f"These experiences have sharpened my skills in {match_str}, "
        f"which I understand are central to this role.\n\n"
        f"{('My academic background in ' + req.education + ' has given me a strong theoretical foundation to complement my practical skills. ') if req.education else ''}"
        f"I am excited about the opportunity to bring this expertise to {req.companyName} and grow alongside your team.\n\n"
        f"Thank you for considering my application. I look forward to discussing how I can contribute.\n\n"
        f"Sincerely,\n{req.candidateName}"
    )


@app.get("/health", tags=["health"])
def health() -> dict[str, str]:
    return {"status": "ok", "service": "placeprep-ai"}


@app.post("/email/generate", tags=["email"])
async def generate_email(req: EmailGenerateRequest) -> StreamingResponse:
    text = build_email_text(req)

    async def stream():
        for char in text:
            yield f"data: {char}\n\n"
            await asyncio.sleep(0.012)
        yield "data: [DONE]\n\n"

    return StreamingResponse(
        stream(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.post("/resume/parse", tags=["resume"])
async def parse_uploaded_resume(request: Request) -> dict:
    content_length = request.headers.get("content-length")
    if content_length and content_length.isdigit() and int(content_length) > MAX_RESUME_BYTES:
        raise HTTPException(status_code=413, detail={"code": "FILE_SIZE_INVALID", "message": "Resume files must be between 1 byte and 5 MB."})
    body_chunks = []
    body_size = 0
    async for chunk in request.stream():
        body_size += len(chunk)
        if body_size > MAX_RESUME_BYTES:
            raise HTTPException(status_code=413, detail={"code": "FILE_SIZE_INVALID", "message": "Resume files must be between 1 byte and 5 MB."})
        body_chunks.append(chunk)
    payload = b"".join(body_chunks)
    filename = request.headers.get("x-file-name", "resume")[:255]
    content_type = request.headers.get("content-type", "").split(";", 1)[0].lower()
    if not payload or len(payload) > MAX_RESUME_BYTES:
        raise HTTPException(status_code=413, detail={"code": "FILE_SIZE_INVALID", "message": "Resume files must be between 1 byte and 5 MB."})

    if content_type == "application/pdf" and payload.startswith(b"%PDF-"):
        text, page_count, used_ocr = extract_pdf(payload)
    elif content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" and payload.startswith(b"PK\x03\x04"):
        text = extract_docx(payload)
        page_count = 0
        used_ocr = False
    else:
        raise HTTPException(status_code=415, detail={"code": "UNSUPPORTED_FILE", "message": "Upload a valid PDF or DOCX resume."})

    parsed = parse_resume(text, page_count, used_ocr)
    parsed["metadata"]["filename"] = filename
    parsed["metadata"]["sha256"] = hashlib.sha256(payload).hexdigest()
    return parsed


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=int(os.getenv("AI_SERVICE_PORT", "8000")))
